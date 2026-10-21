import hashlib
from pathlib import Path
from typing import Iterable

from langchain_core.documents import Document
from langchain_text_splitters import RecursiveCharacterTextSplitter
from langchain_community.document_loaders import PyPDFLoader
from docx import Document as DocxDocument

SUPPORTED = {".pdf", ".txt", ".md", ".docx"}


def _read_text(path: Path) -> str:
    """Read a text file, falling back to other encodings."""
    for encoding in ("utf-8", "utf-8-sig", "cp1252", "latin-1"):
        try:
            return path.read_text(encoding=encoding)
        except UnicodeDecodeError:
            continue
    return path.read_text(encoding="utf-8", errors="ignore")


def _read_docx(path: Path) -> str:
    """Extract paragraphs AND tables from a .docx file."""
    doc = DocxDocument(str(path))
    parts = [p.text.strip() for p in doc.paragraphs if p.text.strip()]

    for table in doc.tables:
        for row in table.rows:
            cells = []
            for cell in row.cells:
                text = cell.text.strip().replace("\n", " ")
                # merged cells repeat the same text; skip consecutive duplicates
                if text and (not cells or cells[-1] != text):
                    cells.append(text)
            if cells:
                parts.append(" | ".join(cells))

    return "\n".join(parts)


def load_file(path: Path) -> list[Document]:
    path = Path(path)
    suffix = path.suffix.lower()
    name = path.name

    if suffix == ".pdf":
        docs = PyPDFLoader(str(path)).load()
    elif suffix in {".txt", ".md"}:
        docs = [Document(page_content=_read_text(path))]
    elif suffix == ".docx":
        docs = [Document(page_content=_read_docx(path))]
    else:
        raise ValueError(f"Unsupported file type: {suffix}")

    # Drop empty pages (e.g. scanned PDFs) and use a clean source name.
    cleaned = []
    for doc in docs:
        if not doc.page_content.strip():
            continue
        doc.metadata["source"] = name
        cleaned.append(doc)

    if not cleaned:
        raise ValueError(
            f"No extractable text in {name}. If it is a scanned PDF, run OCR first."
        )

    return cleaned


def chunk_documents(docs: Iterable[Document]) -> list[Document]:
    splitter = RecursiveCharacterTextSplitter(
        chunk_size=900,
        chunk_overlap=120,
        add_start_index=True,
    )

    chunks = [
        c for c in splitter.split_documents(list(docs))
        if c.page_content.strip()
    ]

    # Deterministic ID per chunk so re-ingesting the same file overwrites
    # instead of duplicating.
    for chunk in chunks:
        key = (
            f"{chunk.metadata.get('source', '')}|"
            f"{chunk.metadata.get('page', '')}|"
            f"{chunk.metadata.get('start_index', '')}|"
            f"{chunk.page_content}"
        )
        chunk.metadata["chunk_id"] = hashlib.sha1(key.encode("utf-8")).hexdigest()

    return chunks